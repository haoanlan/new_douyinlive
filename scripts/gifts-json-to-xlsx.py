#!/usr/bin/env python3
# 把 export-xuanxuan-gifts.js 的 JSON 结果写成多 sheet 的 xlsx
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1]))
out = sys.argv[2]

wb = Workbook()
hdr_font = Font(bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='2F5597')

def sheet(title, rows, cols=None, widths=None):
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(['(无数据)'])
        return ws
    cols = cols or list(rows[0].keys())
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r.get(k) for k in cols])
    ws.freeze_panes = 'A2'
    for i, k in enumerate(cols, 1):
        width = widths.get(k) if widths else None
        if not width:
            vals = [str(r.get(k) or '') for r in rows[:200]] + [str(k)]
            width = min(max(len(v) for v in vals) + 2, 42)
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws

s = data['summary']
sess = data['session']
info = [
    ['场次 ID', sess['id']],
    ['房间', f"{sess['room_id']} / {sess['room_title']}"],
    ['场次开始(北京)', sess['start_time']],
    ['场次结束(北京)', sess['end_time']],
    ['收礼人筛选', data['namePat']],
    ['原始礼物记录条数', s['原始条数']],
    ['连击去重后条数', s['去重条数']],
    ['钻石(连击去重, 仪表盘口径)', s['钻石_去重']],
    ['钻石(原始累计, 未去重)', s['钻石_原始']],
    ['首条礼物时间', s['时间起']],
    ['末条礼物时间', s['时间止']],
    ['本场无收礼人归属记录数', s['无收礼人归属条数']],
    ['本场无收礼人归属钻石(原始)', s['无收礼人归属钻石原始']],
]
ws = wb.active
ws.title = '总览'
for row in info:
    ws.append(row)
ws['A1'].font = Font(bold=True)
ws.column_dimensions['A'].width = 32
ws.column_dimensions['B'].width = 44

sheet('按送礼人', data['userList'], widths={'sec_uid': 30, '送礼人': 24})
sheet('按礼物', data['giftList'], widths={'礼物': 22})
sheet('送礼人x礼物', data['crossList'], widths={'送礼人': 24, '礼物': 22})
sheet('逐条明细', data['records'], widths={'时间': 20, '送礼人': 22, 'sec_uid': 30, '礼物': 20})
sheet('时间线(10分钟)', data['timeline'], widths={'时间段': 20})
sheet('全场断档(>60秒)', data['gaps'], widths={'起': 20, '止': 20})

wb.save(out)
print('wrote', out)
